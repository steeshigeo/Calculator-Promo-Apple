#!/usr/bin/env python3
"""Build the static calculator dataset from the MAP Tech Google Sheet.

The script is deliberately independent of Apps Script.  It downloads the
workbook, normalises the six relevant worksheets, and writes one data.json
file that a static site can load directly.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from urllib.request import Request, urlopen

import openpyxl


SPREADSHEET_ID = "11m3p3cxbqO4VCRS4bR8TrAW8mggTCZEJoKuG2uo_uOs"
DEFAULT_SOURCE_URL = (
    f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/export?format=xlsx"
)

# This list is intentionally explicit.  It keeps banks with no active promo in
# the selector while covering every bank currently present in Promo Berjalan.
CARD_MATRIX = [
    ("BCA", [3, 6, 12, 18, 24]),
    ("BNI", [3, 6, 12, 18, 24]),
    ("Mandiri", [3, 6, 12, 18, 24]),
    ("CIMB", [3, 6, 12, 18, 24]),
    ("BRI CC", [3, 6, 12, 18, 24]),
    ("BRI Samsung Card", [3, 6, 12]),
    ("BRI Debit", [3, 6, 12]),
    ("KB Bank", [3, 6, 12, 18, 24]),
    ("Bank BSI", [3, 6, 12]),
    ("Panin", [3, 6, 12, 18, 24]),
    ("Permata", [3, 6, 12, 18, 24]),
    ("UOB", [3, 6, 12, 18, 24]),
    ("HSBC", [3, 6, 12, 18, 24]),
    ("DBS", [3, 6, 12, 18, 24]),
    ("Maybank", [3, 6, 12, 18, 24]),
    ("OCBC", [3, 6, 12, 18, 24]),
]


def text(value: Any) -> str:
    """Return a stable, whitespace-normalised string for a cell value."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def header(value: Any) -> str:
    return text(value).casefold()


def compact(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def number(value: Any) -> float:
    """Read numeric values in both Indonesian and US-style thousands formats."""
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)

    value = re.sub(r"[^0-9,.-]", "", text(value))
    if not value or value in {"-", ".", ","}:
        return 0.0

    if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+(?:,\d+)?", value):
        value = value.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(?:,\d{3})+(?:\.\d+)?", value):
        value = value.replace(",", "")
    elif "," in value and "." in value:
        # The right-most separator is the decimal separator when both exist.
        if value.rfind(",") > value.rfind("."):
            value = value.replace(".", "").replace(",", ".")
        else:
            value = value.replace(",", "")
    elif value.count(",") == 1 and len(value.rsplit(",", 1)[1]) != 3:
        value = value.replace(",", ".")
    else:
        value = value.replace(",", "")

    try:
        return float(value)
    except ValueError:
        return 0.0


def money(value: Any) -> int:
    return round(number(value))


def row_values(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[tuple[Any, ...]]:
    return list(ws.iter_rows(values_only=True))


def column_index(headers: Iterable[Any], *aliases: str, starts_with: bool = False) -> int:
    """Find a column despite casing and extra spaces in the workbook headers."""
    normal = [header(value) for value in headers]
    aliases_normal = [header(alias) for alias in aliases]
    aliases_compact = [compact(alias) for alias in aliases_normal]
    for index, value in enumerate(normal):
        for alias, alias_compact in zip(aliases_normal, aliases_compact):
            if value == alias or compact(value) == alias_compact:
                return index
            if starts_with and (value.startswith(alias) or compact(value).startswith(alias_compact)):
                return index
    return -1


def value_at(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if 0 <= index < len(row) else None


def require_sheet(workbook: openpyxl.Workbook, name: str) -> openpyxl.worksheet.worksheet.Worksheet:
    if name not in workbook.sheetnames:
        raise ValueError(f"Worksheet '{name}' tidak ditemukan.")
    return workbook[name]


def find_price_header(rows: list[tuple[Any, ...]]) -> tuple[int, tuple[Any, ...], dict[str, int]]:
    for index, row in enumerate(rows[:40]):
        article = column_index(row, "SAP Article")
        description = column_index(row, "SAPDescription", "SAP Description")
        category = column_index(row, "Category")
        normal_price = column_index(row, "Normal Price", starts_with=True)
        promo_price = column_index(row, "Promotion", "Promotion Price", starts_with=True)
        if min(article, description, category, normal_price, promo_price) >= 0:
            return index, row, {
                "article": article,
                "description": description,
                "category": category,
                "normal_price": normal_price,
                "promo_price": promo_price,
                "remarks": column_index(row, "Cash Price Remarks", "Remarks", starts_with=True),
            }
    raise ValueError("Header Price List tidak ditemukan.")


def memory_label(description: str) -> str:
    matches = re.findall(r"(?<![A-Z0-9])(\d+)\s*(TB|GB|T|G)\b", description.upper())
    labels: list[str] = []
    for amount, unit in matches:
        normal_unit = {"G": "GB", "T": "TB"}.get(unit, unit)
        label = f"{amount}{normal_unit}"
        if label not in labels:
            labels.append(label)
    return " / ".join(labels[-2:]) if labels else "Standard"


def catalog_tab(category: str, group: str, description: str) -> str | None:
    signal = f"{category} {group} {description}".casefold()
    if "watch" in signal:
        return "Apple Watch"
    if "iphone" in signal:
        return "iPhone"
    if "ipad" in signal:
        return "iPad"
    if "mac" in signal or "studio display" in signal:
        return "Mac"
    return None


def parse_catalog(workbook: openpyxl.Workbook) -> dict[str, list[dict[str, Any]]]:
    rows = row_values(require_sheet(workbook, "Price List"))
    header_row, _, columns = find_price_header(rows)
    catalog: dict[str, list[dict[str, Any]]] = {
        "iPhone": [],
        "iPad": [],
        "Apple Watch": [],
        "Mac": [],
    }
    current_group = ""

    for row in rows[header_row + 1 :]:
        article = text(value_at(row, columns["article"]))
        description = text(value_at(row, columns["description"]))
        category = text(value_at(row, columns["category"]))

        # Group title rows occasionally carry an accidental value elsewhere;
        # product columns B and C are the reliable discriminator.
        if article and not description and not category:
            current_group = article
            continue
        if not article or not description:
            continue

        tab = catalog_tab(category, current_group, description)
        if not tab:
            continue
        normal_price = money(value_at(row, columns["normal_price"]))
        promo_price = money(value_at(row, columns["promo_price"]))
        if normal_price <= 0 and promo_price <= 0:
            continue
        catalog[tab].append(
            {
                "group": current_group or description,
                "description": description,
                "article": article,
                "normal_price": normal_price,
                "promo_price": promo_price,
                "remarks": text(value_at(row, columns["remarks"])),
                "memory_label": memory_label(description),
            }
        )
    return catalog


def find_priced_header(rows: list[tuple[Any, ...]], label: str) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows[:25]):
        article = column_index(row, "SAP Article")
        description = column_index(row, "SAP Description", "SAPDescription")
        repricing = column_index(row, "Repricing Cash", "Repricing", starts_with=True)
        if min(article, description, repricing) >= 0:
            return index, {
                "brand": column_index(row, "Brand"),
                "article": article,
                "description": description,
                "reference": column_index(row, "Reference SRP", "Reference", starts_with=True),
                "current": column_index(row, "Current SRP Cash", "Current", starts_with=True),
                "repricing": repricing,
            }
    raise ValueError(f"Header {label} tidak ditemukan.")


def parse_protection_or_provider(
    workbook: openpyxl.Workbook, sheet_name: str, is_qoala: bool
) -> list[dict[str, Any]]:
    rows = row_values(require_sheet(workbook, sheet_name))
    header_row, columns = find_priced_header(rows, sheet_name)
    records: list[dict[str, Any]] = []
    for row in rows[header_row + 1 :]:
        article = text(value_at(row, columns["article"]))
        description = text(value_at(row, columns["description"]))
        if not article or not description:
            continue
        record = {
            "brand": text(value_at(row, columns["brand"])),
            "article": article,
            "description": description,
            "reference": money(value_at(row, columns["reference"])),
            "current": money(value_at(row, columns["current"])),
            "repricing": money(value_at(row, columns["repricing"])),
        }
        if is_qoala:
            maximum = re.search(r"\bmax\s*([0-9.,]+)", description, flags=re.IGNORECASE)
            tenor = re.search(r"\b(\d+)\s*(?:bulan|month)\b", description, flags=re.IGNORECASE)
            record["max_device"] = money(maximum.group(1)) if maximum else 0
            record["qoala_tenor"] = int(tenor.group(1)) if tenor else 0
        records.append(record)
    return records


def parse_trade_in(workbook: openpyxl.Workbook) -> list[dict[str, Any]]:
    rows = row_values(require_sheet(workbook, "Trade in"))
    for header_row, row in enumerate(rows[:30]):
        brand = column_index(row, "Brand")
        model = column_index(row, "Model Name", "Model")
        grade_s = column_index(row, "Grade S")
        if min(brand, model, grade_s) >= 0:
            columns = {
                "brand": brand,
                "model": model,
                "S": grade_s,
                "A": column_index(row, "Grade A"),
                "B": column_index(row, "Grade B"),
                "C": column_index(row, "Grade C"),
                "D": column_index(row, "Grade D"),
            }
            break
    else:
        raise ValueError("Header Trade in tidak ditemukan.")

    records: list[dict[str, Any]] = []
    for row in rows[header_row + 1 :]:
        brand_value = text(value_at(row, columns["brand"]))
        model_value = text(value_at(row, columns["model"]))
        if not brand_value or not model_value:
            continue
        records.append(
            {
                "brand": brand_value,
                "model": model_value,
                "grades": {grade: money(value_at(row, index)) for grade, index in columns.items() if grade in "SABCD"},
            }
        )
    return records


def parse_bnpl(workbook: openpyxl.Workbook) -> list[dict[str, Any]]:
    rows = row_values(require_sheet(workbook, "BNPL"))
    for header_row, row in enumerate(rows[:10]):
        financing = column_index(row, "Financing")
        if financing >= 0:
            break
    else:
        raise ValueError("Header BNPL tidak ditemukan.")

    records: list[dict[str, Any]] = []
    for row in rows[header_row + 1 :]:
        name = text(value_at(row, financing))
        if not name:
            continue
        tenors: list[int] = []
        for value in row[financing + 1 :]:
            if isinstance(value, (int, float)) and not isinstance(value, bool) and value > 0:
                tenor = int(value)
                if tenor not in tenors:
                    tenors.append(tenor)
        if tenors:
            records.append({"name": name, "tenors": tenors})
    return records


def parse_promos(workbook: openpyxl.Workbook) -> list[dict[str, Any]]:
    rows = row_values(require_sheet(workbook, "Promo Berjalan"))
    for header_row, row in enumerate(rows[:15]):
        promo = column_index(row, "Promo")
        bank = column_index(row, "Bank")
        scheme = column_index(row, "Scheme")
        minimum = column_index(row, "Minimal Amount", "Minimum Amount")
        maximum = column_index(row, "Maximal Amount", "Maximum Amount")
        discount = column_index(row, "Discount/Cashback", "Discount", starts_with=True)
        if min(promo, bank, scheme, minimum, maximum, discount) >= 0:
            period = column_index(row, "Periode", "Period")
            break
    else:
        raise ValueError("Header Promo Berjalan tidak ditemukan.")

    promos: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows[header_row + 1 :], start=header_row + 2):
        bank_value = text(value_at(row, bank))
        scheme_value = text(value_at(row, scheme))
        if not bank_value or scheme_value not in {"Direct Discount", "CB by Billing"}:
            continue
        promos.append(
            {
                "key": f"promo-{row_number}",
                "promo": text(value_at(row, promo)),
                "period": text(value_at(row, period)),
                "bank": bank_value,
                "scheme": scheme_value,
                "min": money(value_at(row, minimum)),
                "max": money(value_at(row, maximum)),
                "discount": money(value_at(row, discount)),
            }
        )
    return promos


def build_data(input_path: Path, source_label: str) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    data: dict[str, Any] = {
        "catalog": parse_catalog(workbook),
        "qoala": parse_protection_or_provider(workbook, "Qoala Protection", is_qoala=True),
        "providers": parse_protection_or_provider(workbook, "Provider", is_qoala=False),
        "trade_in": parse_trade_in(workbook),
        "bnpl": parse_bnpl(workbook),
        "card_options": [{"name": name, "tenors": tenors} for name, tenors in CARD_MATRIX],
        "promos": parse_promos(workbook),
        "updated_at": datetime.now(timezone.utc).isoformat(),
        "source_file": source_label,
    }
    fingerprint = json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    data["version"] = hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()[:16]
    return data


def download_source(url: str) -> Path:
    request = Request(url, headers={"User-Agent": "MAP-Tech-Calculator-Sync/1.0"})
    with urlopen(request, timeout=60) as response:  # nosec B310 - URL is operator-configured
        content = response.read()
    if not content.startswith(b"PK"):
        raise ValueError("Sumber tidak mengembalikan file XLSX yang valid.")
    handle = tempfile.NamedTemporaryFile(prefix="maptech-source-", suffix=".xlsx", delete=False)
    handle.write(content)
    handle.close()
    return Path(handle.name)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build data.json from the MAP Tech price workbook.")
    parser.add_argument("--input", type=Path, help="Use a local .xlsx file instead of downloading it.")
    parser.add_argument("--output", type=Path, default=Path("data.json"), help="Output JSON path.")
    parser.add_argument(
        "--source-url",
        default=os.environ.get("GOOGLE_SHEETS_XLSX_URL", DEFAULT_SOURCE_URL),
        help="Google Sheets XLSX export URL. Defaults to GOOGLE_SHEETS_XLSX_URL or the configured sheet.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    temporary_file: Path | None = None
    try:
        if args.input:
            input_path = args.input
            source_label = input_path.name
        else:
            temporary_file = download_source(args.source_url)
            input_path = temporary_file
            source_label = "Google Sheets export"
        if not input_path.is_file():
            raise FileNotFoundError(f"File sumber tidak ditemukan: {input_path}")

        data = build_data(input_path, source_label)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        counts = {
            **{name: len(records) for name, records in data["catalog"].items()},
            "qoala": len(data["qoala"]),
            "providers": len(data["providers"]),
            "trade_in": len(data["trade_in"]),
            "bnpl": len(data["bnpl"]),
            "card_options": len(data["card_options"]),
            "promos": len(data["promos"]),
        }
        print(json.dumps({"output": str(args.output), "version": data["version"], "counts": counts}, ensure_ascii=False))
    finally:
        if temporary_file:
            temporary_file.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
