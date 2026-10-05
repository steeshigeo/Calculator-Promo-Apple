#!/usr/bin/env python3
"""
Sync master data Kalkulator Promo Apple  ->  data.json

Alur (sama dengan Dashtp3staff):
    file master (.xlsx)  --download-->  parsing 6 sheet  -->  data.json  -->  index.html

Sumber file diambil dari environment variable ONEDRIVE_URL (diisi dari GitHub Secret
ONEDRIVE_DIRECT_URL). Link yang didukung:
    - direct link .xlsx (OneDrive / SharePoint / hosting lain)
    - link share OneDrive (1drv.ms / onedrive.live.com / SharePoint)  -> dicoba beberapa cara unduh;
      link HARUS dibuka untuk umum ('Siapa saja yang memiliki link' - Dapat melihat)
    - link Google Sheets (docs.google.com/spreadsheets/d/...)  -> otomatis di-export ke xlsx

Pemakaian lokal (tanpa download):
    python sync_onedrive.py --file "Calculator Promo.xlsx"

Port dari Code.gs (Apps Script) dengan perbaikan:
    - header sheet dibaca tanpa peduli huruf besar/kecil & spasi ganda ("Promotion   Price")
    - baris judul model = kolom A terisi, kolom B & C kosong (tahan sel nyasar)
    - label RAM/Storage Mac (16G/512GB, 24GB/1T, 16GB/256-IND, dst.)
    - daftar bank + tenor kartu kredit dibaca dari tab BANK (kolom A = bank, C-G = tenor); CARD_MATRIX hanya cadangan
    - tiap promo diberi end_date (dibaca dari kolom Periode) supaya promo lewat tidak ditawarkan
"""
import argparse
import base64
import calendar
import hashlib
import io
import json
import os
import re
import sys
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urlparse

import requests
from openpyxl import load_workbook

OUT_FILE = "data.json"
WIB = timezone(timedelta(hours=7))

SHEET_NAMES = {
    "price": "Price List",
    "promo": "Promo Berjalan",
    "bnpl": "BNPL",
    "provider": "Provider",
    "qoala": "Qoala Protection",
    "trade": "Trade in",
    "bank": "BANK",  # opsional: daftar bank + tenor kartu kredit (kolom A = bank, kolom C-G = tenor)
}
REQUIRED_SHEETS = ("price", "promo", "bnpl", "provider", "qoala", "trade")

# Tenor cicilan 0% kartu kredit (tidak ada di sheet -> dikelola di sini).
CARD_MATRIX = [
    ("BCA", [3, 6, 12, 18, 24]),
    ("BRI", [3, 6, 12, 18, 24]),
    ("BNI", [3, 6, 12, 18, 24]),
    ("CIMB", [3, 6, 12, 24]),
    ("Bank BSI", [3, 6, 12, 24]),
    ("Panin", [3, 6, 12]),
    ("Mandiri", [3, 6, 12, 18, 24]),
    ("Permata", [3, 6, 12, 18, 24]),
    ("DBS", [3, 6, 12, 18, 24]),
    ("HSBC", [3, 6, 12, 18, 24]),
    ("Danamon", [3, 6, 12, 18, 24]),
    ("UOB", [3, 6, 12, 18, 24]),
    ("Maybank", [3, 6, 12, 24]),
    ("Jenius (BTPN)", [3, 6, 12]),
    ("KB Bank", [3, 6, 12, 18, 24]),
    ("OCBC", [3, 6, 12]),
]
DEFAULT_NEW_CARD_TENORS = [3, 6, 12]

# Promo yang namanya memuat kata-kata ini diabaikan (tidak masuk data.json / deteksi promo bank).
IGNORE_PROMO_KEYWORDS = ("samsung",)


class SheetError(Exception):
    """Struktur sheet tidak sesuai harapan -> data.json lama dipertahankan."""


# --------------------------------------------------------------------------- helpers
def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():  # sel angka (mis. model "5") jangan jadi "5.0"
        return str(int(v))
    return str(v).strip()


def norm_ws(v):
    return re.sub(r"\s+", " ", text(v)).strip()


def hdr(v):
    return norm_ws(v).lower()


def cell(row, i):
    return row[i] if 0 <= i < len(row) else None


def num(v):
    """Angka dari sel. Nilai numerik dipakai apa adanya; teks dibaca format US/Indonesia."""
    if v is None or v == "" or isinstance(v, bool):
        return 0
    if isinstance(v, (int, float)):
        n = v
    else:
        s = re.sub(r"[^0-9.,\-]", "", str(v))
        if not s or s == "-":
            return 0
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", s):  # 16.999.000 / 16.999.000,50
            s = s.replace(".", "").replace(",", ".")
        else:  # 16,999,000 / 16,999,000.50
            s = s.replace(",", "")
        try:
            n = float(s)
        except ValueError:
            return 0
    if n != n or n in (float("inf"), float("-inf")):
        return 0
    n = round(n, 4)  # buang noise floating point (599400.0000000001)
    return int(n) if float(n).is_integer() else n


def normalize_bank(s):
    """Sama persis dengan normalizeBank() di index.html."""
    s = text(s).lower()
    s = re.sub(r"bank\s*", "", s)
    s = re.sub(r"\s+cc$", "", s)
    s = re.sub(r"\s+card$", "", s)
    return re.sub(r"\s+", "", s).strip()


def find_sheet(sheets, name):
    """Cari sheet tanpa peduli huruf besar/kecil & spasi."""
    want = name.strip().lower()
    for k, v in sheets.items():
        if k.strip().lower() == want:
            return v
    return None


# --------------------------------------------------------------------------- download
UA = {"User-Agent": "Mozilla/5.0 (kalkulator-promo-sync)"}


def host_of(u):
    return urlparse(u).netloc.lower()


def add_download(u):
    if "download=1" in u:
        return u
    return u + ("&" if "?" in u else "?") + "download=1"


def candidate_urls(url):
    """Daftar alamat unduh yang dicoba berurutan untuk satu link master."""
    u = url.strip()
    m = re.match(r"https://docs\.google\.com/spreadsheets/d/([\w-]+)", u)
    if m:
        return [f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"]
    host = host_of(u)
    if host == "1drv.ms" or host.endswith("onedrive.live.com"):
        b64 = base64.urlsafe_b64encode(u.encode()).decode().rstrip("=")
        return [f"https://api.onedrive.com/v1.0/shares/u!{b64}/root/content", add_download(u)]
    if host.endswith("sharepoint.com"):
        return [add_download(u)]
    return [u]


def sharing_hint(host):
    if host.endswith("sharepoint.com") or host == "1drv.ms" or host.endswith("onedrive.live.com"):
        return (
            "Penyebab paling umum: link tidak terbuka untuk umum. Di OneDrive buka file master -> Bagikan -> "
            "Pengaturan link -> pilih 'Siapa saja yang memiliki link' dengan izin 'Dapat melihat' "
            "(bukan 'Orang tertentu' / 'Orang di organisasi Anda'), lalu salin link BARU ke secret ONEDRIVE_DIRECT_URL. "
            "Akun kantor/sekolah (SharePoint) sering memblokir link publik oleh admin; bila begitu, pakai link "
            "Google Sheets atau direct link .xlsx dari hosting lain."
        )
    if "google" in host:
        return "Pastikan Google Sheets dibagikan 'Siapa saja yang memiliki link' (Viewer)."
    return "Pastikan link dapat dibuka tanpa login dan langsung mengunduh file .xlsx."


def download(url):
    """Unduh file master. Mencoba beberapa bentuk link; bila semua gagal, beri diagnosis (tanpa membocorkan URL)."""
    s = requests.Session()
    s.headers.update(UA)
    queue = candidate_urls(url)
    tried, resolved, i = [], False, 0
    first_host = host_of(url)
    while i < len(queue):
        target = queue[i]
        i += 1
        try:
            r = s.get(target, timeout=90, allow_redirects=True)
        except requests.RequestException as e:
            tried.append(f"{host_of(target)} -> gagal koneksi ({type(e).__name__})")
            continue
        if r.status_code == 200 and r.content[:2] == b"PK":
            return r.content
        landed = host_of(r.url)
        note = f"{host_of(target)} -> HTTP {r.status_code}"
        if r.status_code == 200:
            note += " (bukan file xlsx, kemungkinan halaman login)"
        if landed != host_of(target):
            note += f", dialihkan ke {landed}"
        tried.append(note)
        # Link pendek 1drv.ms bisa mengarah ke SharePoint (akun kantor): coba alamat tujuan akhirnya.
        if first_host == "1drv.ms" and not resolved:
            resolved = True
            try:
                final = s.get(url.strip(), timeout=60, allow_redirects=True, stream=True).url
                if host_of(final).endswith(("sharepoint.com", "onedrive.live.com")) and add_download(final) not in queue:
                    queue.append(add_download(final))
            except requests.RequestException:
                pass
    raise SheetError(
        "Gagal mengunduh file master. Percobaan: " + "; ".join(tried) + ". " + sharing_hint(first_host)
    )


def read_sheets(raw):
    if raw[:2] != b"PK":
        raise SheetError("File bukan .xlsx. Simpan ulang file master sebagai .xlsx (Excel Workbook).")
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        is_bank = ws.title.strip().lower() == SHEET_NAMES["bank"].lower()
        rows = []
        for row in ws.iter_rows():
            vals = []
            for c in row:
                v = c.value
                # Di tab BANK sel berformat persen (mis. 0%) dibaca sebagai teks "0%", bukan angka 0 (= tidak tersedia).
                if is_bank and isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in (c.number_format or ""):
                    v = f"{v * 100:g}%"
                vals.append(v)
            rows.append(vals)
        sheets[ws.title.strip()] = rows
    return sheets


# --------------------------------------------------------------------------- parsers
def extract_memory_label(description):
    s = norm_ws(description)
    # Mac: spesifikasi dipisah "/" ; RAM & storage ada di dua token berurutan.
    if "/" in s:
        parts = [p.strip() for p in re.sub(r"-IND$", "", s.upper()).split("/")]
        for i in range(len(parts) - 2, -1, -1):
            ram = re.fullmatch(r"(\d+)\s*(?:GB|G)?", parts[i])
            sto = re.fullmatch(r"(\d+)\s*(GB|TB|T)?(?:-IND)?", parts[i + 1])
            if ram and sto:
                n = int(sto.group(1))
                if sto.group(2):
                    unit = "TB" if sto.group(2) == "T" else sto.group(2)
                else:
                    unit = "GB" if n >= 64 else "TB"
                return f"{ram.group(1)}GB / {n}{unit}"
        if not re.search(r"\d+\s*(GB|TB)", s, re.I):
            return s  # tanpa RAM/storage (mis. Studio Display): pakai deskripsi penuh
    m = re.search(r"(\d+\s*GB)\s*[/%x×]\s*(\d+\s*(?:GB|TB))", s, re.I)
    if m:
        return re.sub(r"\s+", "", m.group(1)) + " / " + re.sub(r"\s+", "", m.group(2))
    m = re.search(r"(\d+\s*(?:GB|TB))(?!.*\d+\s*(?:GB|TB))", s, re.I)
    return re.sub(r"\s+", "", m.group(1)) if m else "Standard"


def tab_for(category, description, group):
    if category == "Apple Device - iPhone":
        return "iPhone"
    if category == "Apple Device - iPad":
        return "iPad"
    if "APPLE WATCH" in description.upper() or group.lower().startswith("apple watch"):
        return "Apple Watch"
    return "Mac"


def parse_price_list(rows):
    hi = None
    for r in range(min(len(rows), 15)):
        h = [hdr(v) for v in rows[r]]
        if (
            "sap article" in h and "sapdescription" in h and "category" in h
            and "normal price" in h and any(x.startswith("promotion") for x in h)
        ):
            hi = r
            break
    if hi is None:
        raise SheetError(
            "Header sheet 'Price List' tidak ditemukan. Harus ada kolom: SAP Article, SAPDescription, "
            "Category, Normal Price, Promotion Price (di 15 baris pertama)."
        )
    h = [hdr(v) for v in rows[hi]]
    ia, ib, ic = h.index("sap article"), h.index("sapdescription"), h.index("category")
    inormal = h.index("normal price")
    ipromo = next(i for i, x in enumerate(h) if x.startswith("promotion"))
    ichange = next((i for i, x in enumerate(h) if x.startswith("penurunan")), -1)
    iremarks = h.index("cash price remarks") if "cash price remarks" in h else -1

    catalog = {"iPhone": [], "iPad": [], "Apple Watch": [], "Mac": []}
    current = ""
    for row in rows[hi + 1:]:
        a, b, c = text(cell(row, ia)), text(cell(row, ib)), text(cell(row, ic))
        if a and not b and not c:  # baris judul model
            current = a
            continue
        if not (a and b and c) or not re.match(r"^Apple Device - ", c, re.I):
            continue
        normal = num(cell(row, inormal))
        promo = num(cell(row, ipromo))
        group = norm_ws(current or b)
        catalog[tab_for(c, b, group)].append({
            "article": a,
            "description": b,
            "category": c,
            "group": group,
            "normal_price": normal,
            "promo_price": promo or normal,
            "change": num(cell(row, ichange)) if ichange >= 0 else 0,
            "remarks": text(cell(row, iremarks)) if iremarks >= 0 else "",
            "memory_label": extract_memory_label(b),
        })
    return catalog


def parse_catalog(rows):
    """Sheet Provider & Qoala Protection."""
    if not rows:
        return []
    hi = 0
    for r in range(min(5, len(rows))):
        h = [hdr(v) for v in rows[r]]
        if "sap article" in h and ("sap description" in h or "sapdescription" in h):
            hi = r
            break
    h = [hdr(v) for v in rows[hi]]

    def find(pred):
        return next((i for i, x in enumerate(h) if pred(x)), -1)

    ib = find(lambda x: x == "brand")
    ia = find(lambda x: x == "sap article")
    idesc = find(lambda x: x in ("sap description", "sapdescription"))
    iref = find(lambda x: x.startswith("reference"))
    icur = find(lambda x: x.startswith("current"))
    irep = find(lambda x: "repricing" in x)
    if ia < 0:
        raise SheetError("Kolom 'SAP Article' tidak ditemukan di sheet Provider/Qoala.")
    if irep < 0:
        raise SheetError("Kolom 'Repricing Cash' tidak ditemukan di sheet Provider/Qoala.")

    out = []
    for row in rows[hi + 1:]:
        article = text(cell(row, ia))
        if not article:
            continue
        out.append({
            "brand": text(cell(row, ib)) if ib >= 0 else "",
            "article": article,
            "description": text(cell(row, idesc)) if idesc >= 0 else article,
            "reference": num(cell(row, iref)) if iref >= 0 else 0,
            "current": num(cell(row, icur)) if icur >= 0 else 0,
            "repricing": num(cell(row, irep)),
        })
    return out


def parse_qoala(rows):
    qoala = parse_catalog(rows)
    for q in qoala:
        m = re.search(r"Max\s*([0-9.]+)", q["description"], re.I)
        q["max_device"] = int(m.group(1).replace(".", "")) if m else 0
        m = re.search(r"(\d+)\s*bulan", q["description"], re.I)
        tenor = int(m.group(1)) if m else None
        if not tenor:
            m = re.search(r"\((\d+)-", q["article"])
            tenor = int(m.group(1)) if m else None
        q["qoala_tenor"] = tenor
    return qoala


def parse_bnpl(rows):
    if len(rows) < 2:
        return []
    h = [hdr(v) for v in rows[0]]
    iname = h.index("financing") if "financing" in h else 0
    ibunga = next((i for i, x in enumerate(h) if x.startswith("bunga")), 4)
    out = []
    for row in rows[1:]:
        name = text(cell(row, iname))
        if not name:
            continue
        tenors = [int(n) for n in (num(cell(row, c)) for c in range(iname + 1, ibunga)) if n]
        out.append({"name": name, "tenors": tenors, "interest": text(cell(row, ibunga))})
    return out


def parse_trade_in(rows):
    if len(rows) < 3:
        return []
    hi = None
    for r in range(min(5, len(rows))):
        h = [hdr(v) for v in rows[r]]
        if "brand" in h and "model name" in h:
            hi = r
            break
    if hi is None:
        raise SheetError("Header sheet 'Trade in' tidak ditemukan (butuh kolom BRAND dan MODEL NAME).")
    h = [hdr(v) for v in rows[hi]]
    ibrand, imodel = h.index("brand"), h.index("model name")
    gcols = {g: h.index("grade " + g.lower()) for g in "SABCD" if ("grade " + g.lower()) in h}
    out = []
    for row in rows[hi + 1:]:
        brand, model = text(cell(row, ibrand)), text(cell(row, imodel))
        if not brand or not model:
            continue
        grades = {g: num(cell(row, i)) for g, i in gcols.items() if cell(row, i) not in (None, "")}
        if grades:
            out.append({"brand": brand, "model": model, "grades": grades})
    return out


MONTHS = {
    "januari": 1, "jan": 1, "februari": 2, "feb": 2, "maret": 3, "mar": 3, "april": 4, "apr": 4,
    "mei": 5, "juni": 6, "jun": 6, "juli": 7, "jul": 7, "agustus": 8, "agu": 8, "agt": 8, "ags": 8,
    "september": 9, "sep": 9, "sept": 9, "oktober": 10, "okt": 10, "november": 11, "nov": 11,
    "desember": 12, "des": 12,
}


def period_end(period, today):
    """Tanggal berakhir dari teks Periode (mis. '24 April - 13 September', 'Hingga 31 Desember'). None bila tidak terbaca."""
    words = [w for w in re.findall(r"[a-z]+", text(period).lower()) if w in MONTHS]
    if not words:
        return None
    start_m, end_m = MONTHS[words[0]], MONTHS[words[-1]]
    days = re.findall(r"(\d{1,2})\s*([a-z]+)", text(period).lower())
    days = [(int(d), MONTHS[w]) for d, w in days if w in MONTHS]
    year = today.year
    if start_m > end_m and today.month >= start_m:  # melintasi pergantian tahun (Des -> Jan)
        year += 1
    if days:
        day, month = days[-1]
    else:
        day, month = 31, end_m
    day = min(day, calendar.monthrange(year, month)[1])
    try:
        return date(year, month, day)
    except ValueError:
        return None


def normalize_scheme(raw):
    low = text(raw).lower()
    if "direct" in low or "potong" in low:
        return "Direct Discount"
    if "cashback" in low or "billing" in low or re.match(r"^cb\b", low):
        return "CB by Billing"
    return text(raw)


PROMO_COLS = {
    "promo": ("promo", "nama promo"),
    "period": ("periode", "period"),
    "bank": ("bank",),
    "scheme": ("scheme", "skema"),
    "min": ("minimal amount", "minimum amount", "min amount", "minimal", "min"),
    "max": ("maximal amount", "maksimal amount", "maximum amount", "max amount", "maksimal", "max"),
}


def parse_promos(rows, today, warnings=None):
    hi, idx = None, {}
    for r in range(min(len(rows), 6)):
        h = [hdr(v) for v in rows[r]]
        found = {}
        for key, names in PROMO_COLS.items():
            i = next((h.index(n) for n in names if n in h), -1)
            if i < 0:
                break
            found[key] = i
        else:
            hi, idx = r, found
            break
    if hi is None:
        raise SheetError(
            "Header sheet 'Promo Berjalan' tidak ditemukan. Butuh kolom: Promo, Periode, Bank, Scheme, "
            "Minimal Amount, Maximal Amount, Discount/Cashback."
        )
    h = [hdr(v) for v in rows[hi]]
    idisc = next((i for i, x in enumerate(h) if x.startswith(("discount", "cashback", "diskon"))), -1)
    if idisc < 0:
        raise SheetError("Kolom 'Discount/Cashback' tidak ditemukan di sheet Promo Berjalan.")
    out, carry, unknown = [], "", set()
    for row in rows[hi + 1:]:
        promo, period = text(cell(row, idx["promo"])), text(cell(row, idx["period"])) or carry
        bank, scheme = text(cell(row, idx["bank"])), normalize_scheme(cell(row, idx["scheme"]))
        if text(cell(row, idx["period"])):
            carry = text(cell(row, idx["period"]))
        if not promo or not bank or not scheme:
            continue
        if scheme not in ("Direct Discount", "CB by Billing"):
            unknown.add(scheme)
        end = period_end(period, today)
        out.append({
            "promo": promo, "period": period, "bank": bank, "scheme": scheme,
            "min": num(cell(row, idx["min"])), "max": num(cell(row, idx["max"])), "discount": num(cell(row, idisc)),
            "end_date": end.isoformat() if end else None,
        })
    if unknown and warnings is not None:
        warnings.append(f"Scheme promo tidak dikenali (ditampilkan tanpa potongan): {sorted(unknown)}. Gunakan 'Direct Discount' atau 'CB by Billing'.")
    return out


NEG_MARKS = {"-", "–", "—", "x", "×", "✗", "✕", "tidak", "no", "n/a", "na", "false", "n", "none", "null"}
BANK_HEADERS = {"bank", "nama bank", "banks", "bank name"}
_PURE_TENOR = re.compile(r"(\d{1,2})\s*(x|bln|bulan|months?|mo|kali)?")


def bank_tenor(head, v):
    """Tenor (bulan) yang diwakili sel kolom C-G tab BANK; 0 bila tidak tersedia.

    - Header angka ("3", "12 bulan"): sel berisi tanda apa pun (centang, Y, 0%, 1) = tenor itu tersedia; kosong / "-" / 0 = tidak.
    - Header bukan angka ("Tenor 1"): angka di dalam sel (3, 6, 12 ...) adalah tenornya.
    """
    if v is None or v == "":
        return 0
    s = text(v).lower()
    if s in NEG_MARKS:
        return 0
    isnum = isinstance(v, (int, float)) and not isinstance(v, bool)
    pure = _PURE_TENOR.fullmatch(head.strip().lower())
    if pure:
        return 0 if (isnum and v == 0) else int(pure.group(1))
    if isnum and float(v).is_integer() and 1 <= v <= 60:
        return int(v)
    m = re.fullmatch(r"\D*?(\d{1,2})\s*(x|bln|bulan|months?|mo|kali)?\D*", s)
    if m and not s.endswith("%") and 1 <= int(m.group(1)) <= 60:
        return int(m.group(1))
    loose = re.search(r"(\d{1,2})", head)
    return int(loose.group(1)) if loose else 0


def parse_bank(rows):
    """Tab BANK: kolom A = nama bank, kolom C-G = tenor cicilan kartu kredit."""
    if not rows:
        return []
    hi = 0
    for r in range(min(6, len(rows))):
        if norm_ws(cell(rows[r], 0)).lower() in BANK_HEADERS:
            hi = r
            break
    heads = {c: norm_ws(cell(rows[hi], c)) for c in range(2, 7)}
    out, index = [], {}
    for row in rows[hi + 1:]:
        name = text(cell(row, 0))
        if not name or name.lower() in BANK_HEADERS:
            continue
        tenors = {t for t in (bank_tenor(heads[c], cell(row, c)) for c in range(2, 7)) if t}
        key = normalize_bank(name)
        if key in index:
            index[key]["tenors"] = sorted(set(index[key]["tenors"]) | tenors)
            continue
        rec = {"name": name, "tenors": sorted(tenors)}
        index[key] = rec
        out.append(rec)
    return out


def promo_bank_key(name):
    return normalize_bank(re.sub(r"(?i)debit", "", text(name)))


def build_banks(rows, promos, warnings):
    """card_options (kartu kredit + tenor) dan debit_options dari tab BANK; fallback ke CARD_MATRIX bila tab tidak ada."""
    if rows is None:
        warnings.append("Tab 'BANK' tidak ditemukan -> memakai daftar bank bawaan (CARD_MATRIX). Tambahkan tab BANK agar daftar bank & tenor mengikuti sheet.")
        banks, source = [{"name": n, "tenors": list(t)} for n, t in CARD_MATRIX], "fallback"
    else:
        banks, source = parse_bank(rows), "sheet"
        if not banks:
            raise SheetError("Tab 'BANK' kosong: isi nama bank di kolom A dan tenor di kolom C-G.")
        if not any(b["tenors"] for b in banks):
            warnings.append("Tab BANK: tidak ada tenor yang terbaca dari kolom C-G. Cek header kolom (mis. 3, 6, 12, 18, 24) dan tanda di sel.")
        for b in banks:
            print(f"::notice::BANK {b['name']}: tenor {', '.join(map(str, b['tenors'])) if b['tenors'] else '(tanpa cicilan, hanya debit)'}")
    keys = {normalize_bank(b["name"]) for b in banks}
    missing = sorted({p["bank"] for p in promos if promo_bank_key(p["bank"]) not in keys})
    if missing:
        warnings.append(f"Bank di tab Promo Berjalan yang tidak ada di tab BANK (promonya tidak bisa dipilih): {missing}")
    return banks, [{"name": b["name"]} for b in banks], source


# --------------------------------------------------------------------------- build
def find_label_collisions(catalog):
    """Varian yang harganya beda tapi label RAM/Storage sama -> hanya satu yang bisa dipilih di kalkulator."""
    out = []
    for tab, items in catalog.items():
        seen = {}
        for p in items:
            conn = ""
            if tab == "iPad":
                d = p["description"].lower()
                conn = "C" if re.search(r"wi-?fi\s*\+?\s*cell|cellular", d) else "W"
            seen.setdefault((p["group"], conn, p["memory_label"]), set()).add(p["promo_price"])
        for (group, conn, label), prices in seen.items():
            if len(prices) > 1:
                out.append(f"[{tab}] {group} / {label}{' / ' + conn if conn else ''}: harga {sorted(prices)}")
    return out


def build(sheets, today, source_name, warnings):
    def get(key):
        rows = find_sheet(sheets, SHEET_NAMES[key])
        if rows is None and key in REQUIRED_SHEETS:
            raise SheetError(f"Sheet '{SHEET_NAMES[key]}' tidak ditemukan. Sheet yang ada: {', '.join(sheets)}")
        return rows

    catalog = parse_price_list(get("price"))
    promos_all = parse_promos(get("promo"), today, warnings)
    promos = [p for p in promos_all if not any(k in p["promo"].lower() for k in IGNORE_PROMO_KEYWORDS)]
    ignored = len(promos_all) - len(promos)
    if ignored:
        warnings.append(f"{ignored} promo diabaikan karena nama promo memuat kata {IGNORE_PROMO_KEYWORDS}.")
    banks, debit, bank_source = build_banks(get("bank"), promos, warnings)
    data = {
        "source_file": source_name,
        "catalog": catalog,
        "bnpl": parse_bnpl(get("bnpl")),
        "providers": parse_catalog(get("provider")),
        "qoala": parse_qoala(get("qoala")),
        "trade_in": parse_trade_in(get("trade")),
        "card_options": banks,
        "debit_options": debit,
        "bank_source": bank_source,
        "promos": promos,
    }
    # Pengaman: jangan publish data kosong/rusak.
    for tab, items in catalog.items():
        if not items:
            raise SheetError(f"Katalog '{tab}' kosong setelah parsing. Cek kolom Category di sheet Price List.")
    for key in ("bnpl", "providers", "qoala", "trade_in"):
        if not data[key]:
            raise SheetError(f"Data '{key}' kosong setelah parsing.")
    if not promos:
        warnings.append("Tidak ada promo bank di tab Promo Berjalan (kalkulator menampilkan 'tidak ada promo').")
    collisions = find_label_collisions(catalog)
    if collisions:
        warnings.append(
            f"{len(collisions)} kelompok produk punya >1 varian dengan label RAM/Storage sama tetapi harga berbeda "
            "(mis. warna / Plus / chip). Di kalkulator dipilih lewat dropdown 'Varian / Warna'. "
            "Rinciannya: jalankan lokal dengan --verbose."
        )
        if os.environ.get("VERBOSE"):
            warnings.extend(collisions)
    expired = [p for p in promos if p["end_date"] and p["end_date"] < today.isoformat()]
    if expired:
        warnings.append(
            f"{len(expired)} dari {len(promos)} promo sudah lewat periodenya (disembunyikan otomatis di kalkulator); "
            f"sebaiknya baris-nya dihapus dari sheet Promo Berjalan."
        )
    unreadable = [p for p in promos if not p["end_date"]]
    if unreadable:
        warnings.append(f"{len(unreadable)} promo dengan Periode yang tidak terbaca tanggalnya (tetap ditampilkan).")
    return data


def rp(n):
    return "Rp" + f"{int(n):,}".replace(",", ".")


def catalog_index(data):
    out = {}
    for tab, items in (data.get("catalog") or {}).items():
        for p in items:
            out[(tab, p["article"], p["description"])] = p
    return out


def diff_report(old, new, limit=40):
    """Daftar perubahan harga produk antara data.json lama dan hasil sync ini."""
    a, b = catalog_index(old), catalog_index(new)
    changed, added = [], []
    for k, p in b.items():
        o = a.get(k)
        if o is None:
            added.append(f"[{k[0]}] {p['description']} (baru) promo {rp(p['promo_price'])}, normal {rp(p['normal_price'])}")
            continue
        bits = []
        if o["promo_price"] != p["promo_price"]:
            bits.append(f"promo {rp(o['promo_price'])} -> {rp(p['promo_price'])}")
        if o["normal_price"] != p["normal_price"]:
            bits.append(f"normal {rp(o['normal_price'])} -> {rp(p['normal_price'])}")
        if bits:
            changed.append(f"[{k[0]}] {p['description']}: " + ", ".join(bits))
    removed = [f"[{k[0]}] {p['description']} (dihapus dari sheet)" for k, p in a.items() if k not in b]
    return changed, added, removed


def print_diff(old, new):
    changed, added, removed = diff_report(old, new)
    lines = changed + added + removed
    for ln in lines[:40]:
        print(f"::notice::{ln}")
    if len(lines) > 40:
        print(f"::notice::... dan {len(lines) - 40} perubahan lain")
    print(f"Ringkasan perubahan: {len(changed)} harga berubah, {len(added)} produk baru, {len(removed)} produk dihapus.")
    summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary and lines:
        with open(summary, "a", encoding="utf-8") as f:
            f.write("### Perubahan data kalkulator\n\n" + "\n".join(f"- {ln}" for ln in lines[:200]) + "\n")


def fingerprint(data):
    blob = json.dumps(data, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha1(blob.encode("utf-8")).hexdigest()[:16]


def total_products(data):
    return sum(len(v) for v in data.get("catalog", {}).values())


def main():
    ap = argparse.ArgumentParser(description="Sync master Excel -> data.json")
    ap.add_argument("--file", help="pakai file .xlsx lokal (tanpa download)")
    ap.add_argument("--out", default=OUT_FILE)
    ap.add_argument("--source-name", default=os.environ.get("SOURCE_NAME", "Calculator Promo"))
    ap.add_argument("--force", action="store_true", help="tulis ulang walau data tidak berubah")
    ap.add_argument("--verbose", action="store_true", help="tampilkan rincian varian berlabel sama")
    args = ap.parse_args()
    if args.verbose:
        os.environ["VERBOSE"] = "1"

    now = datetime.now(WIB)
    warnings = []
    try:
        if args.file:
            raw = open(args.file, "rb").read()
        else:
            url = os.environ.get("ONEDRIVE_URL") or os.environ.get("SOURCE_URL")
            if not url:
                raise SheetError("ONEDRIVE_URL kosong. Isi GitHub Secret ONEDRIVE_DIRECT_URL (Settings > Secrets > Actions).")
            raw = download(url)
        data = build(read_sheets(raw), now.date(), args.source_name, warnings)
    except SheetError as e:
        print(f"::error::{e}")
        return 1
    except requests.RequestException as e:
        print(f"::error::Gagal mengunduh file master: {e}")
        return 1

    version = fingerprint(data)
    old = None
    if os.path.exists(args.out):
        try:
            old = json.load(open(args.out, encoding="utf-8"))
        except (OSError, ValueError):
            old = None

    for w in warnings:
        print(f"::warning::{w}")

    if old and total_products(old) and total_products(data) < total_products(old) * 0.5:
        print(f"::error::Jumlah produk turun drastis ({total_products(old)} -> {total_products(data)}). "
              "Dicurigai sheet terpotong/rusak; data.json lama dipertahankan.")
        return 1

    cat = {k: len(v) for k, v in data["catalog"].items()}
    print(f"Terbaca: katalog {cat} | qoala {len(data['qoala'])} | provider {len(data['providers'])} | "
          f"trade-in {len(data['trade_in'])} | bnpl {len(data['bnpl'])} | kartu {len(data['card_options'])} | promo {len(data['promos'])}")

    if old and old.get("version") == version and not args.force:
        print("Tidak ada perubahan data. data.json tidak diubah.")
        return 0

    if old:
        print_diff(old, data)
    data["version"] = version
    data["updated_at"] = now.isoformat(timespec="seconds")
    tmp = args.out + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    os.replace(tmp, args.out)
    print(f"data.json diperbarui (versi {version}, {os.path.getsize(args.out) / 1024:.0f} KB).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
